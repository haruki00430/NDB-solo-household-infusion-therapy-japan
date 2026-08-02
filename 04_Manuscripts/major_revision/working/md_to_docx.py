"""
Minimal Markdown -> DOCX converter for the Major Revision response letter,
cover letter, and Online Resource 1 supplement.

Supports: #/##/### headings, **bold** inline runs, > blockquotes,
--- horizontal rules (rendered as a thin border paragraph), pipe tables,
and plain paragraphs. Intentionally simple; these three documents do not
need anything more elaborate.
"""
import re
import sys
from pathlib import Path

import docx
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

IMAGE_RE = re.compile(r"^!\[[^\]]*\]\(([^)]+)\)$")


BREAK_MARK = "\x00BR\x00"


def add_inline_runs(paragraph, text, base_italic=False):
    for chunk_idx, chunk in enumerate(text.split(BREAK_MARK)):
        if chunk_idx > 0:
            paragraph.add_run().add_break()
        # split on **bold** segments
        parts = re.split(r"(\*\*[^*]+\*\*)", chunk)
        for part in parts:
            if not part:
                continue
            if part.startswith("**") and part.endswith("**"):
                run = paragraph.add_run(part[2:-2])
                run.bold = True
            else:
                run = paragraph.add_run(part)
            if base_italic:
                run.italic = True


def add_horizontal_rule(document):
    p = document.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "999999")
    pBdr.append(bottom)
    pPr.append(pBdr)


def set_row_repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)


def set_row_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    cantSplit = OxmlElement("w:cantSplit")
    trPr.append(cantSplit)


def set_paragraph_keep_next(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    keepNext = OxmlElement("w:keepNext")
    pPr.append(keepNext)


def add_table(document, lines):
    rows = [ln.strip().strip("|").split("|") for ln in lines if ln.strip()]
    rows = [[c.strip() for c in row] for row in rows]
    # drop the markdown separator row (---|---|---)
    rows = [row for row in rows if not all(re.fullmatch(r"-+", c) for c in row)]
    if not rows:
        return
    ncols = len(rows[0])
    table = document.add_table(rows=0, cols=ncols)
    table.style = "Light Grid Accent 1" if "Light Grid Accent 1" in [s.name for s in document.styles] else None
    for i, row in enumerate(rows):
        table_row = table.add_row()
        cells = table_row.cells
        for j, val in enumerate(row[:ncols]):
            cells[j].text = ""
            p = cells[j].paragraphs[0]
            add_inline_runs(p, val)
            if i == 0:
                for run in p.runs:
                    run.bold = True
        # repeat the header row on every page the table spans
        if i == 0:
            set_row_repeat_header(table_row)
        # never split a single row across a page break
        set_row_cant_split(table_row)
        # chain keepNext through every row except the last so the whole
        # table is kept together on one page whenever it fits
        if i < len(rows) - 1:
            for cell in cells:
                for p in cell.paragraphs:
                    set_paragraph_keep_next(p)
    document.add_paragraph()


def convert(md_path: Path, out_path: Path, title: str = None):
    document = docx.Document()
    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    lines = md_path.read_text(encoding="utf-8").splitlines()
    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if stripped == "---":
            add_horizontal_rule(document)
            i += 1
            continue

        if stripped.startswith(("### ", "## ", "# ")):
            if stripped.startswith("### "):
                level, heading_text = 3, stripped[4:]
            elif stripped.startswith("## "):
                level, heading_text = 2, stripped[3:]
            else:
                level, heading_text = 1, stripped[2:]
            heading = document.add_heading(heading_text.strip(), level=level)
            j = i + 1
            while j < n and not lines[j].strip():
                j += 1
            if j < n and lines[j].strip().startswith("|"):
                set_paragraph_keep_next(heading)
            i += 1
            continue

        if stripped.startswith("|"):
            table_lines = []
            while i < n and lines[i].strip().startswith("|"):
                table_lines.append(lines[i])
                i += 1
            add_table(document, table_lines)
            continue

        img_match = IMAGE_RE.match(stripped)
        if img_match:
            img_path = Path(img_match.group(1))
            if not img_path.is_absolute():
                img_path = (md_path.parent / img_path).resolve()
            document.add_picture(str(img_path), width=Inches(5.8))
            i += 1
            continue

        if stripped.startswith(">"):
            quote_lines = []
            while i < n and lines[i].strip().startswith(">"):
                quote_lines.append(lines[i].strip().lstrip(">").strip())
                i += 1
            p = document.add_paragraph()
            p.paragraph_format.left_indent = Pt(24)
            add_inline_runs(p, " ".join(quote_lines), base_italic=True)
            continue

        # plain paragraph: merge consecutive non-blank lines (standard Markdown
        # paragraph semantics) until a blank line, a structural marker, or the
        # start of a new numbered-list item (each "N. " starts its own paragraph).
        # A line ending in a hard-break marker (trailing "  " or "\") forces a
        # visual line break within the same paragraph instead of prose-joining.
        raw_first = line.rstrip("\n")
        hard_break = raw_first.endswith("  ") or raw_first.endswith("\\")
        para_segments = [(stripped, hard_break)]
        i += 1
        while i < n:
            raw_next = lines[i].rstrip("\n")
            nxt = raw_next.strip()
            if not nxt or nxt.startswith(("#", "|", ">", "---")) or re.match(r"^\d+\.\s", nxt):
                break
            nxt_hard_break = raw_next.endswith("  ") or raw_next.endswith("\\")
            para_segments.append((nxt, nxt_hard_break))
            i += 1
        joined = para_segments[0][0]
        prev_hard_break = para_segments[0][1]
        for seg, seg_hard_break in para_segments[1:]:
            if prev_hard_break:
                sep = BREAK_MARK
            else:
                prev_char = joined[-1] if joined else ""
                next_char = seg[0] if seg else ""
                sep = " " if prev_char.isascii() and prev_char.isalnum() and next_char.isascii() and next_char.isalnum() else ""
            joined += sep + seg
            prev_hard_break = seg_hard_break
        p = document.add_paragraph()
        add_inline_runs(p, joined)

    document.save(str(out_path))
    print(f"[OK] {out_path}")


if __name__ == "__main__":
    md_file = Path(sys.argv[1])
    out_file = Path(sys.argv[2])
    convert(md_file, out_file)
