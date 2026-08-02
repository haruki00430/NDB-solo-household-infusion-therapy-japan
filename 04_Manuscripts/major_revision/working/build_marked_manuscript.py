"""
Build the marked-up manuscript via a genuine Word document comparison
(Application.CompareDocuments) between the originally submitted manuscript
and the corrected clean manuscript. This shows word-level insertions
(underlined) and deletions (struck through), which is required to preserve
a full, granular revision history -- a coarser paragraph-level highlight
does not show what specifically changed within a paragraph.

Note: python-docx's `Paragraph.text` does not read `<w:delText>` content
inside `<w:del>` revision runs, so inspecting a tracked-changes document
with plain `.text` will look like text is missing/garbled even though the
document is fully intact. Always verify via rendered PDF (or via a
full-text helper that also reads `<w:delText>`), not via `.text` alone.

Word substitutes its own localized default reviewer name (e.g. "作成者")
in `w:author` attributes even when `Application.UserName` is set before
comparing, so the resulting docx is post-processed to replace it with a
neutral "Author" string.
"""
from pathlib import Path
import os
import zipfile
import shutil

import win32com.client
from docx.oxml.ns import qn

PROJECT_ROOT = Path(r"C:\Users\user\.ag-cursor-common\research_workspace\projects\NDB_Research_Hub\projects\NDB_XXX_heatwave_heatstroke")
ORIGINAL_DOCX = PROJECT_ROOT / "04_Manuscripts" / "submission_package_IJB" / "Japan manuscript_main_IJB_anon.docx"
CLEAN_DOCX = PROJECT_ROOT / "04_Manuscripts" / "major_revision" / "final" / "manuscript_main_IJB_major_revision_clean.docx"
OUT_DOCX = PROJECT_ROOT / "04_Manuscripts" / "major_revision" / "final" / "manuscript_main_IJB_major_revision_marked.docx"

LEAKED_AUTHOR_STRINGS = ["作成者", "Author"]  # scrub any localized default; leave a clean "Author" behind


def run_word_compare():
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    word.UserName = "Author"
    word.UserInitials = "AU"

    doc_orig = word.Documents.Open(str(ORIGINAL_DOCX))
    doc_rev = word.Documents.Open(str(CLEAN_DOCX))

    compared = word.CompareDocuments(
        OriginalDocument=doc_orig,
        RevisedDocument=doc_rev,
        Destination=2,  # wdCompareDestinationNew
        Granularity=1,  # wdGranularityWordLevel
        CompareFormatting=False,
        CompareCaseChanges=True,
        CompareWhitespace=True,
        CompareTables=True,
        CompareHeaders=True,
        CompareFootnotes=True,
        CompareTextboxes=True,
        CompareFields=True,
        CompareComments=False,
        CompareMoves=True,
        RevisedAuthor="Author",
        IgnoreAllComparisonWarnings=True,
    )

    if OUT_DOCX.exists():
        os.remove(OUT_DOCX)
    compared.SaveAs(str(OUT_DOCX), FileFormat=16)  # wdFormatDocumentDefault (.docx)
    pages = compared.ComputeStatistics(2)

    compared.Close(False)
    doc_orig.Close(False)
    doc_rev.Close(False)
    word.Quit()
    return pages


def scrub_author_metadata():
    tmp = str(OUT_DOCX) + ".tmp"
    with zipfile.ZipFile(OUT_DOCX, "r") as zin:
        names = zin.namelist()
        data = {n: zin.read(n) for n in names}

    target = "word/document.xml"
    xml = data[target].decode("utf-8")
    xml = xml.replace('w:author="作成者"', 'w:author="Author"')
    data[target] = xml.encode("utf-8")

    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for n in names:
            zout.writestr(n, data[n])
    shutil.move(tmp, OUT_DOCX)


def strip_comments():
    """
    The originally submitted base docx has Reviewer 1's inline review comments
    embedded in it. CompareComments=False does not stop CompareDocuments from
    carrying those pre-existing comment anchors into unchanged spans of the
    output, so every regeneration re-introduces them. All 35 items are already
    answered in the response letter, so the marked manuscript should show only
    tracked changes, not a partial subset of reviewer comment balloons.
    """
    import re

    tmp = str(OUT_DOCX) + ".tmp"
    with zipfile.ZipFile(OUT_DOCX, "r") as zin:
        names = zin.namelist()
        data = {n: zin.read(n) for n in names}

    comment_parts = {"word/comments.xml", "word/commentsExtended.xml", "word/commentsIds.xml"}

    doc_xml = data["word/document.xml"].decode("utf-8")
    doc_xml = re.sub(r'<w:commentRangeStart[^/]*/>', '', doc_xml)
    doc_xml = re.sub(r'<w:commentRangeEnd[^/]*/>', '', doc_xml)
    doc_xml = re.sub(r'<w:r>(?:(?!<w:r>|</w:r>).)*?<w:commentReference[^/]*/>.*?</w:r>', '', doc_xml, flags=re.DOTALL)
    doc_xml = re.sub(r'<w:commentReference[^/]*/>', '', doc_xml)
    data["word/document.xml"] = doc_xml.encode("utf-8")

    for part in comment_parts:
        data.pop(part, None)
    names = [n for n in names if n not in comment_parts]

    rels_path = "word/_rels/document.xml.rels"
    rels_xml = data[rels_path].decode("utf-8")
    rels_xml = re.sub(r'<Relationship[^>]*Target="comments[^"]*"[^>]*/>', '', rels_xml)
    data[rels_path] = rels_xml.encode("utf-8")

    ct_path = "[Content_Types].xml"
    ct_xml = data[ct_path].decode("utf-8")
    ct_xml = re.sub(r'<Override PartName="/word/comments[^"]*"[^>]*/>', '', ct_xml)
    data[ct_path] = ct_xml.encode("utf-8")

    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for n in names:
            zout.writestr(n, data[n])
    shutil.move(tmp, OUT_DOCX)


def accept_and_highlight_revised_tables():
    """
    Word's CompareDocuments diffs replaced tables cell-by-cell at the same
    row/column position, not as a semantic "this table is gone, that one is
    new" change. For Table 1 (revised in place) and Table 2 (different row
    structure entirely -- O-A/O-B/O-C/O-D vs the original model list), this
    produces unreadable interleaved insertions/deletions within cells and
    a leftover 9-column table shape. Tables with no counterpart in the new
    manuscript (fully superseded old sensitivity tables) render fine as-is
    (entirely struck through, clearly legible as "removed"), so only tables
    that contain at least one insertion are treated: deletions are dropped,
    insertions are unwrapped (accepted), and the whole table is highlighted
    yellow to flag it as revised-in-full rather than diffed cell-by-cell.
    The response letter states explicitly that Table 1 and Table 2 were
    replaced in full.
    """
    import docx
    from docx.enum.text import WD_COLOR_INDEX

    document = docx.Document(OUT_DOCX)
    ins_tag = qn("w:ins")
    del_tag = qn("w:del")
    move_from_tag = qn("w:moveFrom")
    move_to_tag = qn("w:moveTo")
    range_marker_tags = [
        qn("w:moveFromRangeStart"), qn("w:moveFromRangeEnd"),
        qn("w:moveToRangeStart"), qn("w:moveToRangeEnd"),
    ]

    for table in document.tables:
        tbl_el = table._tbl
        has_insertion = tbl_el.find(f".//{ins_tag}") is not None
        if not has_insertion:
            continue  # fully-superseded old table: leave as a clean strikethrough deletion

        # drop text deleted or moved away from this position
        for tag in (del_tag, move_from_tag):
            for el in list(tbl_el.iter(tag)):
                el.getparent().remove(el)
        # accept text inserted here or moved to this position (unwrap, keep content)
        for tag in (ins_tag, move_to_tag):
            for el in list(tbl_el.iter(tag)):
                parent = el.getparent()
                idx = list(parent).index(el)
                for i, child in enumerate(list(el)):
                    parent.insert(idx + i, child)
                parent.remove(el)
        # remove now-empty move range bookmarks
        for tag in range_marker_tags:
            for el in list(tbl_el.iter(tag)):
                el.getparent().remove(el)

        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.font.highlight_color = WD_COLOR_INDEX.YELLOW

    document.save(OUT_DOCX)


def main():
    pages = run_word_compare()
    scrub_author_metadata()
    strip_comments()
    accept_and_highlight_revised_tables()
    print(f"[OK] Word-native tracked-changes manuscript written to {OUT_DOCX} ({pages} pages)")


if __name__ == "__main__":
    main()
