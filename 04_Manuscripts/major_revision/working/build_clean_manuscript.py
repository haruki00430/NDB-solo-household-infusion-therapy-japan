"""
Build the clean revised manuscript DOCX from manuscript_draft.md, using the
original submitted manuscript_main_IJB.docx as the style/template base
(same fonts, margins, Title/Heading 1/Heading 2/Body Text/Normal styles).

No tracked changes, no comments. Continuous line numbering is added via
direct sectPr XML manipulation (python-docx has no high-level API for this).
"""
import re
from pathlib import Path

import docx
import pandas as pd
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, Inches

PROJECT_ROOT = Path(r"C:\Users\user\.ag-cursor-common\research_workspace\projects\NDB_Research_Hub\projects\NDB_XXX_heatwave_heatstroke")
BASE_DOCX = PROJECT_ROOT / "04_Manuscripts" / "submission_package_IJB" / "manuscript_main_IJB.docx"
DRAFT_MD = PROJECT_ROOT / "04_Manuscripts" / "major_revision" / "working" / "manuscript_draft.md"
OUT_DOCX = PROJECT_ROOT / "04_Manuscripts" / "major_revision" / "final" / "manuscript_main_IJB_major_revision_clean.docx"
TABLE1_CSV = PROJECT_ROOT / "03_Analysis" / "results" / "major_revision" / "tables" / "table1_descriptive_statistics.csv"
TABLE2_CSV = PROJECT_ROOT / "03_Analysis" / "results" / "major_revision" / "tables" / "table2_original_exposure_hierarchy_final.csv"
FIGURE1_PNG = PROJECT_ROOT / "03_Analysis" / "results" / "major_revision" / "figures" / "figure_original_exposure_coefficients.png"

MAJOR_HEADINGS = {
    "Abstract", "Introduction", "Materials and methods", "Results", "Discussion",
    "Conclusions", "Acknowledgments", "Conflicts of Interest", "Data Availability",
    "Author Contributions", "Funding",
    "Declaration of Generative AI and AI-Assisted Technologies in the Manuscript Preparation Process",
    "References", "Tables and Figures",
}


def add_continuous_line_numbers(document):
    sectPr = document.sections[0]._sectPr
    lnNumType = OxmlElement('w:lnNumType')
    lnNumType.set(qn('w:countBy'), '1')
    lnNumType.set(qn('w:restart'), 'continuous')
    sectPr.append(lnNumType)


def add_inline_runs(paragraph, text):
    """Parse **bold** markers into runs."""
    parts = re.split(r'(\*\*[^*]+\*\*)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part:
            paragraph.add_run(part)


def insert_table_from_csv(document, csv_path, caption, col_widths_in=None, font_pt=9):
    df = pd.read_csv(csv_path)
    p = document.add_paragraph()
    p.add_run(caption).bold = True
    table = document.add_table(rows=1, cols=len(df.columns))
    table.autofit = False
    tbl = table._tbl
    tblBorders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '4')
        el.set(qn('w:color'), '000000')
        tblBorders.append(el)
    tblPr = tbl.tblPr
    tblPr.append(tblBorders)
    tblLayout = OxmlElement('w:tblLayout')
    tblLayout.set(qn('w:type'), 'fixed')
    tblPr.append(tblLayout)

    def set_cell_font(cell, text, bold=False):
        cell.text = ''
        para = cell.paragraphs[0]
        run = para.add_run(str(text))
        run.font.size = Pt(font_pt)
        run.bold = bold

    hdr_cells = table.rows[0].cells
    for i, col in enumerate(df.columns):
        set_cell_font(hdr_cells[i], col, bold=True)
    for _, row in df.iterrows():
        cells = table.add_row().cells
        for i, val in enumerate(row):
            set_cell_font(cells[i], val)

    if col_widths_in:
        for row in table.rows:
            for i, w in enumerate(col_widths_in):
                row.cells[i].width = Inches(w)
        for i, w in enumerate(col_widths_in):
            table.columns[i].width = Inches(w)

    document.add_paragraph()


def main():
    document = docx.Document(BASE_DOCX)

    # remove all existing body paragraphs and tables
    body = document.element.body
    for child in list(body):
        if child.tag != qn('w:sectPr'):
            body.remove(child)

    lines = DRAFT_MD.read_text(encoding='utf-8').splitlines()

    i = 0
    first_body_para_done = {}
    stop_parsing = False
    while i < len(lines):
        if stop_parsing:
            break
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue
        if line.startswith('# '):
            title_text = line[2:].strip()
            p = document.add_paragraph(style='Title')
            p.add_run(title_text)
        elif line.strip() == '---':
            pass
        elif line.startswith('## '):
            heading_text = line[3:].strip()
            if heading_text == 'Tables and Figures':
                stop_parsing = True  # remaining content is inserted natively below
                continue
            p = document.add_paragraph(style='Heading 1')
            p.add_run(heading_text)
        elif line.startswith('### '):
            heading_text = line[4:].strip()
            p = document.add_paragraph(style='Heading 2')
            p.add_run(heading_text)
        elif line.startswith('**Table') or line.startswith('**Fig') or line.startswith('**Online Resource'):
            p = document.add_paragraph(style='Body Text')
            add_inline_runs(p, line)
        elif re.match(r'^[¹²]', line) or line.startswith('*Corresponding') or (i < 6 and ('Department of' in line or 'Radiation Medical' in line or line.strip().endswith('Ohira¹,²'))):
            p = document.add_paragraph(style='Body Text')
            add_inline_runs(p, line)
        elif line.startswith('- '):
            p = document.add_paragraph(style='Body Text')
            add_inline_runs(p, '\u2022 ' + line[2:])
        elif line.startswith('> '):
            p = document.add_paragraph(style='Body Text')
            add_inline_runs(p, line[2:])
        else:
            p = document.add_paragraph(style='Body Text')
            add_inline_runs(p, line)
        i += 1

    # Native Table 1 (8 columns: Variable, Unit, N, Mean, SD, Min, Median, Max)
    document.add_paragraph(style='Heading 2').add_run('Table 1. Descriptive Statistics of Prefecture-Level Variables (N = 47)')
    insert_table_from_csv(
        document, TABLE1_CSV, '',
        col_widths_in=[2.3, 1.0, 0.4, 0.65, 0.65, 0.65, 0.65, 0.65],
    )

    # Native Table 2 (10 columns)
    document.add_paragraph(style='Heading 2').add_run('Table 2. Original-Exposure Hierarchical Models (O-A to O-D)')
    insert_table_from_csv(
        document, TABLE2_CSV, '',
        col_widths_in=[0.45, 1.35, 0.3, 0.95, 0.95, 0.45, 0.5, 0.5, 0.4, 0.5],
        font_pt=8,
    )

    # Figure 1
    document.add_paragraph(style='Heading 2').add_run('Figure 1')
    fig_p = document.add_paragraph()
    run = fig_p.add_run()
    run.add_picture(str(FIGURE1_PNG), width=Inches(5.5))
    cap_p = document.add_paragraph(style='Body Text')
    cap_p.add_run(
        'Fig. 1 Coefficient plot for the primary exposure (percentage of all households with an '
        'older adult living alone) across models O-A through O-D, showing unstandardized HC3 95% '
        'confidence intervals and a zero reference line.'
    )

    add_continuous_line_numbers(document)

    OUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUT_DOCX)
    print(f"[OK] clean manuscript written to {OUT_DOCX}")


if __name__ == "__main__":
    main()
