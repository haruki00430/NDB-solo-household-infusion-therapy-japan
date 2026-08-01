"""
IJB 投稿向け表記統一: Keywords 同期、Online Resource、Fig. キャプション。

Usage:
    python apply_ijb_terminology.py
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

sys.stdout.reconfigure(encoding="utf-8")

PACKAGE = Path(__file__).resolve().parent
MANUSCRIPT = PACKAGE / "manuscript_main_IJB.docx"
TITLE_PAGE = PACKAGE / "title_page_IJB.docx"

KEYWORDS_MANUSCRIPT = (
    "heat-health surveillance; social isolation; elderly solo household; "
    "dehydration; ecological study; Japan"
)

SI_REPLACEMENTS = [
    ("Supplementary Figure S1a", "Online Resource 1, panel a"),
    ("Supplementary Figure S1b", "Online Resource 1, panel b"),
    ("Supplementary Figure S1.", "Online Resource 1."),
    ("Supplementary Figure S1", "Online Resource 1"),
]


def replace_in_paragraph(paragraph: Paragraph, replacements: list[tuple[str, str]]) -> bool:
    """段落テキストを置換する。"""
    text = paragraph.text
    new_text = text
    for old, new in replacements:
        new_text = new_text.replace(old, new)
    if new_text == text:
        return False
    for run in paragraph.runs:
        run.text = ""
    if paragraph.runs:
        paragraph.runs[0].text = new_text
    else:
        paragraph.add_run(new_text)
    return True


def replace_in_document(doc: Document, replacements: list[tuple[str, str]]) -> int:
    """全段落・表セルを置換し、変更件数を返す。"""
    count = 0
    for para in doc.paragraphs:
        if replace_in_paragraph(para, replacements):
            count += 1
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    if replace_in_paragraph(para, replacements):
                        count += 1
    return count


def set_fig_caption(paragraph: Paragraph, fig_num: int, caption_body: str | None = None) -> None:
    """IJB 形式の図キャプション（Fig. N 太字、番号後に句点なし）を設定する。"""
    paragraph.clear()
    label = paragraph.add_run(f"Fig. {fig_num}")
    label.bold = True
    if caption_body:
        paragraph.add_run(f" {caption_body}")


def fix_figure_sections(doc: Document) -> None:
    """末尾の Figure Legends / プレースホルダを IJB 形式に更新する。"""
    fig_captions = {
        1: (
            "Scatter plot showing the association between elderly solo household rate (%) "
            "and infusion therapy utilization (per 100,000 population) across 47 Japanese "
            "prefectures, stratified by the median elderly solo household rate (12.3%). "
            "Solid line represents the linear regression fit within each stratum, with shaded "
            "area indicating 95% confidence interval. Top 5 and bottom 5 prefectures by "
            "infusion therapy rate are labeled."
        ),
        2: (
            "Scatter plot showing the association between elderly solo household rate (%) "
            "and infusion therapy utilization (per 100,000 population) across all 47 Japanese "
            "prefectures. Solid line represents the simple linear regression fit "
            "(β = 723.37, p = 0.0001, R² = 0.282), with shaded area indicating "
            "95% confidence interval."
        ),
    }

    for para in doc.paragraphs:
        t = para.text.strip()
        m = re.match(r"^Figure (\d)\.\s*(.+)$", t, re.DOTALL)
        if m:
            num = int(m.group(1))
            body = m.group(2).strip()
            set_fig_caption(para, num, body)
            continue
        if t in ("Figure 1", "Figure 2"):
            num = int(t.split()[-1])
            set_fig_caption(para, num)
            continue
        if t == "Online Resource 1":
            para.clear()
            run = para.add_run("Online Resource 1")
            run.bold = True
            continue
        if t.startswith("Online Resource 1. "):
            body = t[len("Online Resource 1. ") :]
            para.clear()
            label = para.add_run("Online Resource 1")
            label.bold = True
            para.add_run(f" {body}")
            continue


def update_title_page() -> None:
    """タイトルページの Keywords を本文と同期する。"""
    doc = Document(TITLE_PAGE)
    paras = doc.paragraphs
    for i, para in enumerate(paras):
        if para.text.strip() == "Keywords" and i + 1 < len(paras):
            paras[i + 1].text = KEYWORDS_MANUSCRIPT
            break
    for para in doc.paragraphs:
        if para.text.startswith("Tables: 2"):
            para.text = "Tables: 2 | Figures: 2 main + 1 online resource (ESM_1)"
    doc.save(TITLE_PAGE)


def main() -> None:
    doc = Document(MANUSCRIPT)

    # Keywords（本文）
    for para in doc.paragraphs:
        if para.text.startswith("Keywords:"):
            para.text = f"Keywords: {KEYWORDS_MANUSCRIPT}"

    n = replace_in_document(doc, SI_REPLACEMENTS)
    fix_figure_sections(doc)

    doc.save(MANUSCRIPT)
    update_title_page()

    print(f"Updated: {MANUSCRIPT.name}")
    print(f"  Supplementary → Online Resource replacements: {n} paragraph/cell edits")
    print(f"Updated: {TITLE_PAGE.name} (Keywords synced)")


if __name__ == "__main__":
    main()
